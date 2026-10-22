"""Banks and NBFCs (6.1): the lender template, by rules only. A lender's year from its results filings (or an Excel
export), defaults, a projection of the balance sheet and P&L of a lender, residual-income valuation (P/B, P/ABV,
justified P/B, implied RoE), forensic checks and an Excel workbook with live formulas. Never imports ai.

NIM here is NII / average total assets (not interest-earning assets: the filings don't split them), so the RoA tree
reads in one unit: NIM + fee & other income - opex - credit cost - tax = RoA, and RoE = RoA x leverage."""
import io
import math
import statistics

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import model
from model import SCENARIOS, _div, expand, fy, labels

DRIVERS = [  # key, label, kind (as model.DRIVERS)
    ("adv_growth", "Loan growth (advances)", "pct"),
    ("nim", "NIM (NII % avg total assets)", "pct"),
    ("oi_pct", "Fee & other income % avg assets", "pct"),
    ("cost_income", "Cost-to-income", "pct"),
    ("credit_cost", "Credit cost (provisions % avg advances)", "pct"),
    ("tax_rate", "Tax rate", "pct"),
    ("payout", "Dividend payout", "pct"),
    ("adv_assets", "Advances % total assets", "pct"),
    ("rwa_density", "RWA % total assets", "pct"),
    ("cet1_min", "CET1 floor (equity raised below it)", "pct"),
]
KEYS = [k for k, _, _ in DRIVERS]
# bull / bear around base, as seed() in web/src/lib.ts: loan growth +/- 3 pp, NIM +/- 15 bps, credit cost -/+ 25 bps
SHIFT = {"adv_growth": 0.03, "nim": 0.0015, "credit_cost": -0.0025}

LAYOUT = [  # (label, key, kind); the Statements page and the Excel Model sheet
    ("ASSUMPTIONS", None, "head"),
    *[(label, key, kind) for key, label, kind in DRIVERS],
    ("", None, None),
    ("PROFIT & LOSS", None, "head"),
    ("Interest earned", "int_earned", "num"), ("Interest expended", "int_expended", "num"),
    ("Net interest income (NII)", "nii", "num"), ("Fee & other income", "other_income", "num"),
    ("Net total income", "total_income", "num"), ("Operating expenses", "opex", "num"),
    ("Pre-provision operating profit (PPOP)", "ppop", "num"), ("Provisions & contingencies", "provisions", "num"),
    ("Profit before tax", "pbt", "num"), ("Tax", "tax", "num"), ("Net profit (PAT)", "pat", "num"),
    ("Dividend", "dividend", "num"),
    ("", None, None),
    ("BALANCE SHEET", None, "head"),
    ("Share capital", "share_cap", "num"), ("Reserves", "reserves", "num"), ("Net worth", "equity", "num"),
    ("Deposits", "deposits", "num"), ("Borrowings", "borrowings", "num"), ("Other liabilities & provisions", "other_liab", "num"),
    ("Total liabilities", "total_liab", "num"),
    ("Cash & balances with RBI / banks", "cash", "num"), ("Investments", "investments", "num"),
    ("Advances (loans)", "advances", "num"), ("Fixed assets", "fixed_assets", "num"), ("Other assets", "other_assets", "num"),
    ("Total assets", "total_assets", "num"),
    ("", None, None),
    ("CAPITAL", None, "head"),
    ("Risk-weighted assets (proxy)", "rwa", "num"), ("CET1 ratio", "cet1", "pct"), ("Equity raised", "raised", "num"),
    ("", None, None),
    ("ASSET QUALITY (reported)", None, "head"),
    ("Gross NPA", "gnpa", "num"), ("Net NPA", "nnpa", "num"), ("GNPA %", "gnpa_pct", "pct"), ("NNPA %", "nnpa_pct", "pct"),
    ("Provision coverage (PCR)", "pcr", "pct"),
    ("", None, None),
    ("RATIOS", None, "head"),
    ("Yield on assets", "yield", "pct"), ("Cost of funds", "cof", "pct"),
    ("Return on assets (RoA)", "roa", "pct"), ("Return on equity (RoE)", "roe", "pct"),
    ("Leverage (avg assets / avg equity)", "leverage", "x"), ("Credit-deposit ratio", "cd_ratio", "pct"),
    ("", None, None),
    ("CHECKS (should be 0)", None, "head"),
    ("Assets - liabilities", "chk_bs", "chk"),
]
REQUIRED = ("int_earned", "int_expended", "other_income", "opex", "pbt", "tax", "pat", "share_cap", "reserves",
            "total_assets", "advances")  # numbers in every lender year (server.Company checks them)
EXPORT_NOTE = ("Older years from the export: provisions sit inside operating expenses and advances include other assets "
               "(an Excel export also keeps deposits inside borrowings), so credit cost and the provision lines start at "
               "the first filed year")


# ---------- a lender's year ----------

def filing_row(x):
    """An annual results filing's facts (model.parse_xbrl, with LENDER_TAGS) -> one lender year, or None when it has
    no balance sheet or profit lines. Banks file interest expended and provisions; NBFCs (Ind AS) file finance costs
    and impairment, and their fee income sits in revenue from operations. Exceptional items go into other income
    (as model.filing_row does), so PBT = NII + other income - opex - provisions always ties: opex is the remainder."""
    g = lambda k: x.get(k) or 0.0  # noqa: E731
    if x.get("int_earned") is None or x.get("pbt") is None or x.get("pat") is None or not x.get("share_cap") \
            or not x.get("advances"):
        return None
    bank = "int_expended" in x
    int_exp = g("int_expended") if bank else g("interest")
    oi = g("other_income") + g("exceptional") + (0.0 if bank else g("revenue") - x["int_earned"])
    prov = g("provisions")
    reserves = (x["reserves"] if "reserves" in x else x["other_equity"] if "other_equity" in x
                else g("equity_parent") - x["share_cap"])
    fixed = g("fixed_assets") if bank else sum(g(k) for k in ("ppe", "cwip", "intangibles", "intangibles_wip",
                                                                 "investment_property", "goodwill"))
    cash, inv, adv = g("cash") + g("bank_balances"), g("investments_lender"), x["advances"]
    total = x.get("total_assets") or x.get("equity_and_liabilities") or (
        cash + inv + adv + fixed + g("other_assets_lender"))
    nii = x["int_earned"] - int_exp
    tax = x["tax"] if "tax" in x else g("current_tax") + g("deferred_tax")
    r = {"int_earned": x["int_earned"], "int_expended": int_exp, "other_income": oi, "provisions": prov,
         "opex": nii + oi - prov - x["pbt"], "pbt": x["pbt"], "tax": tax, "pat": x["pat"],
         "dividend": abs(g("dividends_paid")), "share_cap": x["share_cap"], "reserves": reserves,
         "deposits": g("deposits"), "borrowings": g("borrowings_lender") + g("debt_securities") + g("subordinated"),
         "cash": cash, "investments": inv, "advances": adv, "fixed_assets": fixed, "total_assets": total}
    r["other_assets"] = total - cash - inv - adv - fixed
    r["other_liab"] = total - x["share_cap"] - reserves - r["deposits"] - r["borrowings"]
    # asset quality and capital are filed only by banks, standalone; a 0 (consolidated files them so) is not filed
    for k in ("gnpa", "nnpa", "gnpa_pct", "nnpa_pct", "cet1", "at1"):
        r[k] = x.get(k) or None
    return r


def export_row(h):
    """A year of model.parse_export's Data Sheet lines for a lender: interest earned is the export's sales, its
    interest the interest expended. The export doesn't split provisions from expenses, deposits from borrowings or
    advances from other assets (EXPORT_NOTE): opex holds provisions, borrowings deposits, advances other assets."""
    total = h["total_assets"]
    r = {"int_earned": h["revenue"], "int_expended": h["interest"], "other_income": h["other_income"], "provisions": None,
         "pbt": h["pbt"], "tax": h["tax"], "pat": h["pat"], "dividend": h["dividend"], "share_cap": h["share_cap"],
         "reserves": h["reserves"], "deposits": None, "borrowings": h["borrowings"], "other_liab": h["other_liab"],
         "cash": h["cash"], "investments": h["investments"], "advances": h["other_assets"] - h["cash"],
         "fixed_assets": h["net_block"] + h["cwip"], "other_assets": 0.0, "total_assets": total,
         "gnpa": None, "nnpa": None, "gnpa_pct": None, "nnpa_pct": None, "cet1": None, "at1": None}
    r["opex"] = r["int_earned"] - r["int_expended"] + r["other_income"] - r["pbt"]  # incl. provisions
    return r


def derive(rows):
    """Lender years (filing_row / export_row) -> history rows with NII, net total income, PPOP, net worth, total
    liabilities and PCR."""
    out = []
    for r in rows:
        h = dict(r)
        h["nii"] = h["int_earned"] - h["int_expended"]
        h["total_income"] = h["nii"] + h["other_income"]
        h["ppop"] = None if h["provisions"] is None else h["total_income"] - h["opex"]
        h["equity"] = h["share_cap"] + h["reserves"]
        h["total_liab"] = h["equity"] + (h["deposits"] or 0.0) + h["borrowings"] + h["other_liab"]
        h["pcr"] = 1 - h["nnpa"] / h["gnpa"] if h.get("gnpa") and h.get("nnpa") is not None else None
        out.append(h)
    return out


def is_lender_export(hist):
    """An Excel export that model._not_lender refused and that looks like a bank or NBFC: receivables and inventory
    under 5% of revenue and interest over a fifth of it. An insurer (premiums, little interest) doesn't, and stays
    refused (model.Lender)."""
    # ponytail: banks run interest at 45-65% of revenue, NBFCs 30-45% (Bajaj Finance FY2026: 31%); 20% leaves room
    rev = [h for h in hist if h["revenue"]]
    med = lambda f: statistics.median(f(h) for h in rev)  # noqa: E731
    return bool(rev) and med(lambda h: (h["receivables"] + h["inventory"]) / h["revenue"]) < 0.05 and \
        med(lambda h: h["interest"] / h["revenue"]) > 0.2


def from_export(co):
    """parse_export's company (before its lender check) -> a lender company."""
    return {**co, "kind": "lender", "hist": derive([export_row(h) for h in co["hist"]]), "quarters": None}


def from_filings(name, annual, market=None, least=2):
    """A lender company from annual results filings (one basis): one year per filing with a balance sheet, oldest
    first -> (company, skipped [{year, why}]). ValueError with fewer than `least` years."""
    years, skipped = {}, []
    for f in annual:
        end = model._iso(f.get("period_end") or "")
        if not end or f.get("kind") != "annual":
            continue
        r = filing_row(f["facts"])
        if r is None:
            skipped.append({"year": fy(end), "why": "the filing has no balance sheet or no lender lines"})
        else:
            years[end] = r
    dates = sorted(years)
    if len(dates) < max(least, 1):
        raise ValueError(("No annual results filing in XBRL has a lender's balance sheet" if not dates else
                          "Only one year of annual results with a balance sheet was filed in XBRL")
                         + "; the model needs at least two years.")
    m = market or {}
    return ({"name": name, "kind": "lender", "dates": dates, "hist": derive([years[d] for d in dates]), "quarters": None,
             "market": {"price": m.get("price"), "shares": m.get("shares"), "face_value": m.get("face_value"),
                        "price_hist": [None] * len(dates), "shares_hist": [None] * len(dates)}}, skipped)


def merge(filed, other):
    """Filings win; the export (a lender company) fills only older years. -> (company, filled years, mismatches
    [{year, line, filing, export}] over 1% on the years both have)."""
    have = {d.date() for d in filed["dates"]}
    add = [(d, h) for d, h in zip(other["dates"], other["hist"]) if d.date() not in have and d < filed["dates"][0]]
    mism = []
    for d, h in zip(other["dates"], other["hist"]):
        if d.date() in have:
            f = filed["hist"][[x.date() for x in filed["dates"]].index(d.date())]
            for k, line in (("pat", "Net profit"), ("total_assets", "Total assets"), ("int_earned", "Interest earned")):
                if abs(f[k] - h[k]) > model.MERGE_TOL * max(abs(f[k]), 1.0):
                    mism.append({"year": fy(d), "line": line, "filing": f[k], "export": h[k]})
    co = {**filed, "dates": [d for d, _ in add] + filed["dates"], "hist": [h for _, h in add] + filed["hist"]}
    m = co.get("market") or {}
    co["market"] = {**m, "price_hist": [None] * len(co["dates"]), "shares_hist": [None] * len(co["dates"])}
    return co, [fy(d) for d, _ in add], mism


def data_checks(co):
    """Assets = liabilities, PBT - tax = net profit (5%: minorities in a consolidated year), 12-month years."""
    issues = []
    for d, h in zip(co["dates"], co["hist"]):
        if abs(h["total_assets"] - h["total_liab"]) > model.TOL:
            issues.append(f"{fy(d)}: assets - liabilities off by {h['total_assets'] - h['total_liab']:,.1f} cr")
        if abs(h["pbt"] - h["tax"] - h["pat"]) > max(model.TOL, 0.05 * abs(h["pat"])):
            issues.append(f"{fy(d)}: PBT - tax vs net profit off by {h['pbt'] - h['tax'] - h['pat']:,.1f} cr "
                          "(minorities or associates, normal for consolidated)")
    d = co["dates"]
    issues += [f"{fy(b)}: period is {(b - a).days} days (year-end change?), its growth/ratios are distorted"
               for a, b in zip(d, d[1:]) if not 330 <= (b - a).days <= 400]
    return issues


# ---------- ratios and defaults ----------

def _avg(h, p, k):
    return (h[k] + p[k]) / 2 if p and p.get(k) is not None and h.get(k) is not None else None


def ratios(hist):
    """Per reported year: the drivers as they were (DRIVERS keys) and the RoA tree. Averages are opening + closing
    over 2, so the first year has none of the flow-over-balance ratios."""
    out = []
    for i, h in enumerate(hist):
        p = hist[i - 1] if i else None
        a, adv, eq = _avg(h, p, "total_assets"), _avg(h, p, "advances"), _avg(h, p, "equity")
        fund = lambda r: (r["deposits"] or 0.0) + r["borrowings"]  # noqa: E731
        f = (fund(h) + fund(p)) / 2 if p else None
        out.append({
            "adv_growth": _div(h["advances"], p["advances"]) - 1 if p and p["advances"] else None,
            "nim": _div(h["nii"], a), "oi_pct": _div(h["other_income"], a),
            "cost_income": _div(h["opex"], h["total_income"]),
            "credit_cost": None if h["provisions"] is None else _div(h["provisions"], adv),
            "tax_rate": _div(h["tax"], h["pbt"]), "payout": _div(h["dividend"], h["pat"]),
            "adv_assets": _div(h["advances"], h["total_assets"]),
            "rwa_density": _div(h["equity"], h["cet1"] * h["total_assets"]) if h.get("cet1") else None,
            "cet1_min": None,
            "yield": _div(h["int_earned"], a), "cof": _div(h["int_expended"], f),
            "opex_assets": _div(h["opex"], a), "prov_assets": None if h["provisions"] is None else _div(h["provisions"], a),
            "tax_assets": _div(h["tax"], a), "roa": _div(h["pat"], a), "roe": _div(h["pat"], eq),
            "leverage": _div(a, eq), "cd_ratio": _div(h["advances"], h["deposits"]) if h["deposits"] else None,
            "equity_assets": _div(h["equity"], h["total_assets"]),
        })
    return out


CET1_FLOOR = (0.12, 0.15)  # bank (CET1 filed): RBI's 8% (5.5% + 2.5% buffer) plus a 4 pp management cushion;
# NBFC (none filed): RBI's 15% CRAR, read against equity / advances
INSURER = "Insurers (life and general) aren’t supported: premiums, reserves and embedded value need their own model."


def defaults(hist, years=3, exclude=(), labels_=None):
    """Rule defaults: loan growth = the median of the last 3 years' growth (a merger year doesn't set it); NIM, fee
    income, cost-to-income, credit cost, tax and payout = ratio of sums over the last 3 years (years with provisions
    for credit cost and cost-to-income: an export year holds provisions in opex); advances % assets and RWA density
    from the last year, so the first projected year starts where the books end."""
    keep = [i for i in range(1, len(hist)) if not labels_ or labels_[i] not in exclude]
    w = keep[-years:] or list(range(1, len(hist)))
    if not w:
        raise ValueError("Needs at least two reported years.")
    r = ratios(hist)
    ros = lambda pairs: model._ros(pairs)  # noqa: E731
    pairs = lambda num, den: [(num(i), den(i)) for i in w]  # noqa: E731
    avg_a = lambda i: (hist[i]["total_assets"] + hist[i - 1]["total_assets"]) / 2  # noqa: E731
    avg_adv = lambda i: (hist[i]["advances"] + hist[i - 1]["advances"]) / 2  # noqa: E731
    full = [i for i in w if hist[i]["provisions"] is not None] or w
    g = [r[i]["adv_growth"] for i in w if r[i]["adv_growth"] is not None]
    last = hist[-1]
    bank = last.get("cet1") is not None
    return {
        "adv_growth": statistics.median(g) if g else 0.0,
        "nim": ros(pairs(lambda i: hist[i]["nii"], avg_a)),
        "oi_pct": ros(pairs(lambda i: hist[i]["other_income"], avg_a)),
        "cost_income": ros([(hist[i]["opex"], hist[i]["total_income"]) for i in full]),
        "credit_cost": ros([(hist[i]["provisions"] or 0.0, avg_adv(i)) for i in full]),
        "tax_rate": ros([(hist[i]["tax"], hist[i]["pbt"]) for i in w if hist[i]["pbt"] > 0]),
        "payout": ros(pairs(lambda i: hist[i]["dividend"] or 0.0, lambda i: hist[i]["pat"] if hist[i]["pat"] > 0 else 0)),
        "adv_assets": last["advances"] / last["total_assets"],
        # CET1 capital ~ net worth (ponytail: it is net worth less intangibles and deductions, so this density is
        # an effective one that makes the first projected ratio read the reported CET1); NBFC: loans at 100% weight
        "rwa_density": last["equity"] / (last["cet1"] * last["total_assets"]) if bank
        else last["advances"] / last["total_assets"],
        "cet1_min": CET1_FLOOR[0] if bank else CET1_FLOOR[1],
    }


# ---------- projection ----------

def _shares(last):
    """The funding and asset mix held from the last reported year: deposits and borrowings as shares of what isn't
    equity; cash and investments as shares of total assets."""
    fund = last["total_assets"] - last["equity"]
    return {"dep": (last["deposits"] or 0.0) / fund if fund else 0.0, "bor": last["borrowings"] / fund if fund else 0.0,
            "inv": last["investments"] / last["total_assets"], "cash": last["cash"] / last["total_assets"]}


def project(hist, a, n=5):
    """The lender's balance sheet and P&L, n years. a: {driver: float or [n floats]}. Total assets follow advances
    (advances % assets); NII and fee income run on average total assets, provisions on average advances; equity
    rolls forward with PAT less dividend, and when CET1 (net worth / RWA) would fall under the floor, equity is
    raised to the floor at year end (shown, and in the RI valuation as a negative dividend). Liabilities fill what
    equity doesn't, split as the last year's. No circularity: nothing in the P&L reads this year's equity."""
    at = lambda k, t: a[k][t] if isinstance(a[k], list) else a[k]  # noqa: E731
    p, mix, out = hist[-1], _shares(hist[-1]), []
    for t in range(n):
        x = {k: at(k, t) for k in KEYS}
        c = {"int_earned": None, "int_expended": None, "gnpa": None, "nnpa": None, "gnpa_pct": None, "nnpa_pct": None,
             "pcr": None, **x}
        c["advances"] = p["advances"] * (1 + x["adv_growth"])
        c["total_assets"] = c["advances"] / x["adv_assets"]
        avg_a = (p["total_assets"] + c["total_assets"]) / 2
        c["nii"] = x["nim"] * avg_a
        c["other_income"] = x["oi_pct"] * avg_a
        c["total_income"] = c["nii"] + c["other_income"]
        c["opex"] = x["cost_income"] * c["total_income"]
        c["ppop"] = c["total_income"] - c["opex"]
        c["provisions"] = x["credit_cost"] * (p["advances"] + c["advances"]) / 2
        c["pbt"] = c["ppop"] - c["provisions"]
        c["tax"] = max(0.0, c["pbt"] * x["tax_rate"])
        c["pat"] = c["pbt"] - c["tax"]
        c["dividend"] = max(0.0, c["pat"] * x["payout"])
        c["rwa"] = x["rwa_density"] * c["total_assets"]
        c["raised"] = max(0.0, x["cet1_min"] * c["rwa"] - (p["equity"] + c["pat"] - c["dividend"]))
        c["share_cap"] = p["share_cap"]
        c["reserves"] = p["reserves"] + c["pat"] - c["dividend"] + c["raised"]
        c["equity"] = c["share_cap"] + c["reserves"]
        c["cet1"] = c["equity"] / c["rwa"] if c["rwa"] > 0 else None
        fund = c["total_assets"] - c["equity"]
        c["deposits"], c["borrowings"] = fund * mix["dep"], fund * mix["bor"]
        c["other_liab"] = fund - c["deposits"] - c["borrowings"]
        c["total_liab"] = c["equity"] + c["deposits"] + c["borrowings"] + c["other_liab"]
        c["investments"], c["cash"] = c["total_assets"] * mix["inv"], c["total_assets"] * mix["cash"]
        c["fixed_assets"] = p["fixed_assets"]
        c["other_assets"] = c["total_assets"] - c["advances"] - c["investments"] - c["cash"] - c["fixed_assets"]
        avg_e = (p["equity"] + c["equity"]) / 2
        c["yield"] = c["cof"] = None
        c["roa"] = c["pat"] / avg_a if avg_a > 0 else None
        c["roe"] = c["pat"] / avg_e if avg_e > 1e-6 else None
        c["leverage"] = avg_a / avg_e if avg_e > 1e-6 else None
        c["cd_ratio"] = c["advances"] / c["deposits"] if c["deposits"] > 1e-6 else None
        c["chk_bs"] = c["total_assets"] - c["total_liab"]
        out.append(c)
        p = c
    return out


# ---------- valuation: residual income ----------

GRID_KE = (-0.02, -0.01, 0.0, 0.01, 0.02)  # sensitivity rows: Ke + step
GRID_ROE = (-0.04, -0.02, 0.0, 0.02, 0.04)  # columns: steady-state RoE + step
ROE_SCALE = [round(0.04 + 0.02 * i, 2) for i in range(14)]  # 4% to 30%, for the implied-RoE ruler
IMPACT = (("adv_growth", "Loan growth", 0.02), ("nim", "NIM", 0.002), ("cost_income", "Cost-to-income", 0.02),
          ("credit_cost", "Credit cost", 0.0025))


def _ri(rows, b0, ke, g, roe=None):
    """Residual income equity value (Rs crore) of projected rows from book b0, or None when Ke doesn't clear g.
    RI_t = PAT_t - Ke x opening book; terminal RI_(n+1) = (RoE - Ke) x closing book, growing at g, where RoE is the
    steady-state return (default: the last year's PAT on its opening book). -> {ri, pv, pv_sum, roe, tv, pv_tv, value}."""
    if not model._clears(ke, g):
        return None
    books = [b0] + [r["equity"] for r in rows]
    ri = [r["pat"] - ke * books[t] for t, r in enumerate(rows)]
    pv = [x / (1 + ke) ** (t + 1) for t, x in enumerate(ri)]
    n = len(rows)
    roe = rows[-1]["pat"] / books[-2] if roe is None and books[-2] > 0 else roe
    if roe is None:
        return None
    tv = (roe - ke) * books[-1] / (ke - g)
    pv_tv = tv / (1 + ke) ** n
    return {"ri": ri, "pv": pv, "pv_sum": sum(pv), "roe": roe, "tv": tv, "pv_tv": pv_tv, "value": b0 + sum(pv) + pv_tv}


def valuation(projs, hist, val, market, inputs=None, note=False):
    """Residual-income valuation of each scenario, from the last reported book. No FCFF DCF: a lender's debt is its
    raw material, so free cash flow to the firm means nothing; equity is valued on book and the return it earns over
    Ke. val: model.VALUATION-shaped (rf, erp, beta, g, wacc_override = a typed Ke, price, shares). Adds P/B and P/ABV
    at CMP, the justified P/B (RoE - g) / (Ke - g), the market-implied steady-state RoE (and Ke), the Ke x RoE grid
    and, with inputs and note=True, what moves value. Mirrors the Valuation sheet in to_excel()."""
    v, last = {**model.VALUATION, **(val or {})}, hist[-1]
    market = market or {}
    price = model._pos(v["price"] if v["price"] is not None else market.get("price"))
    shares = model._pos(v["shares"] if v["shares"] is not None else market.get("shares"))
    ke_build = v["rf"] + v["beta"] * v["erp"]
    ke = v["wacc_override"] if v["wacc_override"] is not None else ke_build
    g, b0 = v["g"], last["equity"]
    error = None
    if not math.isfinite(ke) or ke <= 0:
        error = "Those inputs give no cost of equity. Check the beta, rates and share price."
    elif not model._clears(ke, g):
        error = (f"Ke ({model._pct(ke)}) must be at least 0.5 pp above terminal growth ({model._pct(g)}). "
                 "Lower the growth or raise Ke.")
    ps = lambda x: x / shares if x is not None and shares else None  # noqa: E731
    bvps, nnpa = ps(b0), last.get("nnpa")
    abvps = ps(b0 - nnpa) if nnpa is not None else None
    pb = price / bvps if price and bvps and bvps > 0 else None
    out = {}
    for s, rows in projs.items():
        x = None if error else _ri(rows, b0, ke, g)
        val_ = x and x["value"]
        per = ps(val_)
        roe = x["roe"] if x else (rows[-1]["pat"] / rows[-2]["equity"] if len(rows) > 1 and rows[-2]["equity"] > 0 else None)
        jpb = (roe - g) / (ke - g) if roe is not None and not error else None
        eps = lambda r: ps(r["pat"])  # noqa: E731
        out[s] = {**(x or dict.fromkeys(("ri", "pv", "pv_sum", "tv", "pv_tv"))), "roe_t": roe, "equity": val_,
                  "per_share": per, "upside": per / price - 1 if per is not None and price else None,
                  "error": None if error or per is None or per > 0 else
                  "Equity value comes out at or below zero: the forecast RoE is far below Ke. Check credit cost and NIM.",
                  "justified_pb": jpb, "fv_jpb": jpb * bvps if jpb is not None and bvps else None,
                  "tv_share": x["pv_tv"] / x["value"] if x and x["value"] > 0 else None,
                  "pb_value": per / bvps if per and bvps and bvps > 0 else None,
                  "pe_price": price / eps(rows[0]) if price and eps(rows[0]) and eps(rows[0]) > 0 else None,
                  "pe_value": per / eps(rows[0]) if per and eps(rows[0]) and eps(rows[0]) > 0 else None,
                  "grid": [[ps((_ri(rows, b0, ke + dk, g, roe + dr) or {}).get("value")) if roe is not None and not error
                            else None for dr in GRID_ROE] for dk in GRID_KE],
                  "fcff": [], "fix": None}
    implied_roe = g + pb * (ke - g) if pb is not None and not error else None
    base_roe = out["base"]["roe_t"] if "base" in out else None
    nb = None
    if note and inputs and "base" in inputs and not error:
        nb = {"impact": impact(hist, inputs["base"], b0, ke, g, shares),
              "roe_scale": [[r, ps(b0 * (r - g) / (ke - g))] for r in ROE_SCALE]}
    eps_last = ps(last["pat"])
    return {"kind": "lender", "note": nb, "inputs": v, "price": price, "shares": shares, "ke": ke, "ke_build": ke_build,
            "g": g, "error": error, "book": b0, "bvps": bvps, "abvps": abvps, "nnpa": nnpa,
            "pb": pb, "pabv": price / abvps if price and abvps and abvps > 0 else None,
            "pe": price / eps_last if price and eps_last and eps_last > 0 else None,
            "implied_roe": implied_roe,
            "implied_ke": g + (base_roe - g) / pb if pb and base_roe is not None and not error else None,
            # what the company bar's EV reads (lib.ts capOf): a lender has no EV, so market cap alone
            "borrowings": 0.0, "cash": 0.0, "investments": 0.0, "wacc": ke, "kd": None,
            "steps": {"ke": GRID_KE, "roe": GRID_ROE}, "scenarios": out}


def impact(hist, a, b0, ke, g, shares):
    """Each driver one step down and up, the others at base: value per share either way (the biggest swing first)."""
    n = len(a["nim"])
    out = []
    for k, label, step in IMPACT:
        vals = []
        for d in (-step, step):
            b = {**a, k: [max(0.0, x + d) for x in a[k]]}
            x = _ri(project(hist, b, n), b0, ke, g)
            vals.append(x["value"] / shares if x and shares else None)
        out.append({"key": k, "label": label, "step": step, "down": vals[0], "up": vals[1],
                    "swing": abs(vals[1] - vals[0]) if None not in vals else None})
    return sorted(out, key=lambda r: -(r["swing"] or 0))


# ---------- forensic checks ----------

GNPA_UP = (0.01, 0.005, 0.0025)  # GNPA % up over 2 years: 1 pp high, 50 bps medium, 25 bps low
GNPA_LEVEL = (0.06, 0.03)  # ... or its level: above 6% high, 3% medium
PCR_DOWN = (0.10, 0.05)  # provision coverage down over 2 years
PCR_LEVEL = (0.40, 0.55)  # ... or under 40% high, 55% medium (below: less below is worse)
CC_SPIKE = ((2.0, 0.005), (1.5, 0.0025))  # credit cost against the median of the 3 years before: x and +bps
NIM_DOWN = (0.005, 0.0025, 0.0015)  # NIM down over 2 years
SPREAD_DOWN = (0.005, 0.0025)  # yield on assets less cost of funds, down over 2 years
CD_HIGH = (1.05, 0.95)  # credit-deposit ratio (banks): loans running past deposits
CET1_LOW = (0.09, 0.11, 0.125)  # CET1: under 9% high, 11% medium, 12.5% low
OI_SHARE = (0.50, 0.40)  # fee & other income / net total income: treasury gains and one-offs make profit less steady


def forensic(hist, labels_, notes=None):
    """The lender's forensic radar, in model.forensic's shape. Rules only, on the reported years."""
    out, n, r = [], len(hist), ratios(hist)
    y, last = labels_[-1], hist[-1]

    def add(key, title, sev, value, threshold, years, fact, detail, rule, needs=None):
        out.append({"key": key, "title": title, "severity": sev, "value": value, "threshold": threshold,
                    "years": years, "fact": fact, "detail": detail, "rule": rule, "needs": needs})

    pc = model._pct
    bps = lambda x: f"{x * 1e4:+,.0f} bps"  # noqa: E731
    # GNPA: its level and its rise
    rule = (f"GNPA % now against 2 years before: up {GNPA_UP[2] * 1e4:.0f} bps low, {GNPA_UP[1] * 1e4:.0f} bps medium, "
            f"{GNPA_UP[0] * 1e4:.0f} bps high; or above {pc(GNPA_LEVEL[1])} medium, {pc(GNPA_LEVEL[0])} high")
    have = [(labels_[i], h["gnpa_pct"]) for i, h in enumerate(hist) if h.get("gnpa_pct") is not None]
    if len(have) < 3:
        add("gnpa", "Gross NPAs rising", "not_checked", None, None, [x for x, _ in have],
            "GNPA not filed" if not have else "Needs 3 years of GNPA",
            "Banks file GNPA in their results; NBFCs don't in XBRL, and the annual report's asset-quality note isn't read for it yet.",
            rule)
    else:
        (y0, a), (y1, b) = have[-3], have[-1]
        sev = model._worst(model._sev(b - a, GNPA_UP), model._sev(b, GNPA_LEVEL))
        add("gnpa", "Gross NPAs rising", sev, b, GNPA_LEVEL[1], [x for x, _ in have[-3:]],
            f"GNPA {pc(b)} in {y1} ({bps(b - a)} vs {y0})",
            f"Gross NPAs were {pc(b)} of advances in {y1}, against {pc(a)} in {y0}.", rule)
    # PCR
    rule = (f"Provision coverage (1 - NNPA / GNPA) against 2 years before: down {PCR_DOWN[1]:.0%} medium, "
            f"{PCR_DOWN[0]:.0%} high; or under {PCR_LEVEL[1]:.0%} medium, {PCR_LEVEL[0]:.0%} high")
    have = [(labels_[i], h["pcr"]) for i, h in enumerate(hist) if h.get("pcr") is not None]
    if len(have) < 3:
        add("pcr", "Provision coverage falling", "not_checked", None, None, [x for x, _ in have],
            "PCR not filed" if not have else "Needs 3 years of GNPA and NNPA",
            "Needs gross and net NPAs, which banks file in their results.", rule)
    else:
        (y0, a), (y1, b) = have[-3], have[-1]
        sev = model._worst(model._sev(a - b, PCR_DOWN), model._sev(b, PCR_LEVEL, below=True))
        add("pcr", "Provision coverage falling", sev, b, PCR_LEVEL[1], [x for x, _ in have[-3:]],
            f"PCR {pc(b)} in {y1} ({pc(a)} in {y0})",
            f"Provisions covered {pc(b)} of gross NPAs in {y1}, against {pc(a)} in {y0}.", rule)
    # credit cost spike (the filings carry no slippage line, so provisions stand in for it)
    rule = ("Credit cost (provisions / average advances) against the median of the 3 years before: 1.5x and +25 bps "
            "medium, 2x and +50 bps high. Slippages aren't in the filings; this is their shadow in the P&L")
    cc = [(labels_[i], x["credit_cost"]) for i, x in enumerate(r) if x["credit_cost"] is not None]
    if len(cc) < 4:
        add("credit_cost", "Credit cost spike", "not_checked", None, None, [x for x, _ in cc], "Needs 4 years of provisions",
            "Needs four years with provisions split from operating expenses (filed years; an export doesn't split them).", rule)
    else:
        med, (y1, b) = statistics.median(x for _, x in cc[-4:-1]), cc[-1]
        sev = next((s for s, (m, d) in zip(("high", "medium"), CC_SPIKE) if b >= m * med and b - med >= d), "ok")
        add("credit_cost", "Credit cost spike", sev, b, None, [x for x, _ in cc[-4:]],
            f"Credit cost {pc(b)} in {y1} (median {pc(med)})",
            f"Provisions ran at {pc(b)} of average advances in {y1}, against a median of {pc(med)} in the three years "
            "before.", rule)
    # NIM compression, and the yield / cost-of-funds spread behind it
    nims = [(labels_[i], x["nim"]) for i, x in enumerate(r) if x["nim"] is not None]
    rule = (f"NIM (NII / average assets) against 2 years before: down {NIM_DOWN[2] * 1e4:.0f} bps low, "
            f"{NIM_DOWN[1] * 1e4:.0f} medium, {NIM_DOWN[0] * 1e4:.0f} high")
    if len(nims) < 3:
        add("nim", "NIM compression", "not_checked", None, None, [x for x, _ in nims], "Needs 4 years",
            "Needs four reported years (NIM uses average assets).", rule)
    else:
        (y0, a), (y1, b) = nims[-3], nims[-1]
        add("nim", "NIM compression", model._sev(a - b, NIM_DOWN), b, None, [x for x, _ in nims[-3:]],
            f"NIM {pc(b)} in {y1} ({bps(b - a)} vs {y0})",
            f"NII was {pc(b)} of average assets in {y1}, against {pc(a)} in {y0}.", rule)
    sp = [(labels_[i], x["yield"] - x["cof"]) for i, x in enumerate(r) if x["yield"] is not None and x["cof"] is not None]
    rule = (f"Yield on assets less cost of funds (interest expended / average deposits + borrowings), against 2 years "
            f"before: down {SPREAD_DOWN[1] * 1e4:.0f} bps medium, {SPREAD_DOWN[0] * 1e4:.0f} bps high")
    if len(sp) < 3:
        add("spread", "Yield vs cost of funds", "not_checked", None, None, [x for x, _ in sp], "Needs 4 years",
            "Needs four reported years with interest expended.", rule)
    else:
        (y0, a), (y1, b) = sp[-3], sp[-1]
        i1, i0 = len(r) - 1, labels_.index(y0)
        add("spread", "Yield vs cost of funds", model._sev(a - b, SPREAD_DOWN), b, None, [x for x, _ in sp[-3:]],
            f"Spread {pc(b)} in {y1} ({bps(b - a)} vs {y0})",
            f"Yield on assets {pc(r[i1]['yield'])} against cost of funds {pc(r[i1]['cof'])} in {y1}; "
            f"{pc(r[i0]['yield'])} against {pc(r[i0]['cof'])} in {y0}.", rule)
    # loans outrunning deposits (banks)
    rule = f"Advances / deposits: above {CD_HIGH[1]:.0%} medium, {CD_HIGH[0]:.0%} high (deposit-funded lenders only)"
    if not last["deposits"] or last["deposits"] < 0.5 * (last["total_assets"] - last["equity"]):
        add("cd_ratio", "Loans outrunning deposits", "not_checked", None, None, [y], "Not deposit-funded",
            "Deposits are under half of funding (an NBFC or a wholesale-funded lender), so the ratio doesn't apply.", rule)
    else:
        cd = last["advances"] / last["deposits"]
        add("cd_ratio", "Loans outrunning deposits", model._sev(cd, CD_HIGH), cd, CD_HIGH[1], [y],
            f"Credit-deposit ratio {pc(cd)} in {y}",
            f"Advances were {pc(cd)} of deposits in {y}.", rule)
    # capital
    rule = f"CET1 ratio: under {pc(CET1_LOW[2])} low, {pc(CET1_LOW[1])} medium, {pc(CET1_LOW[0])} high (RBI minimum 8% incl. buffer)"
    if last.get("cet1") is None:
        add("capital", "Thin capital", "not_checked", None, None, [], "CET1 not filed",
            "Banks file CET1 in their results; NBFCs file capital adequacy (CRAR) only in the annual report, which isn't read for it yet.", rule)
    else:
        c1 = last["cet1"]
        add("capital", "Thin capital", model._sev(c1, CET1_LOW, below=True), c1, CET1_LOW[1], [y], f"CET1 {pc(c1)} in {y}",
            f"Common equity was {pc(c1)} of risk-weighted assets in {y}"
            + (f"; AT1 bonds add {pc(last['at1'])} (they can be written down, as Yes Bank's were in 2020)." if last.get("at1")
               else "."), rule)
    # earnings mix
    rule = f"Fee & other income / net total income: above {OI_SHARE[1]:.0%} medium, {OI_SHARE[0]:.0%} high"
    s = last["other_income"] / last["total_income"] if last["total_income"] > 0 else None
    add("oi_share", "Profit leaning on other income", "not_checked" if s is None else model._sev(s, OI_SHARE), s,
        OI_SHARE[1], [y], "No net income" if s is None else f"Other income {pc(s)} of net total income in {y}",
        "" if s is None else f"Fee and other income (incl. treasury gains and exceptional items) were {pc(s)} of net "
        f"total income in {y}.", rule)
    # related parties, from an approved annual-report note (as model.forensic does, on total income)
    yr, nt = model._note(notes, "related")
    rule = f"Related-party transactions / net total income: above {model.NOTE_SHARE[1]:.0%} medium, {model.NOTE_SHARE[0]:.0%} high"
    if nt and isinstance(nt.get("total"), (int, float)) and last["total_income"] > 0:
        sh = nt["total"] / last["total_income"]
        add("related", "Related-party transactions", model._sev(sh, model.NOTE_SHARE), sh, model.NOTE_SHARE[1], [yr],
            f"{pc(sh)} of net total income ({yr})", f"Related-party transactions of {model._cr(nt['total'])} in {yr}.", rule)
    else:
        add("related", "Related-party transactions", "not_checked", None, None, [], "Needs the annual report",
            "Attach an annual report and read its related-party note.", rule, "annual report")
    return out


def key_numbers(hist, labels_):
    """The note's key numbers: the last year against the median of the 3 before (model.key_numbers' shape)."""
    r, y = ratios(hist), labels_[-1]
    out = []
    for k, label in (("nim", "NIM"), ("cost_income", "Cost-to-income"), ("credit_cost", "Credit cost"),
                     ("roa", "RoA"), ("roe", "RoE"), ("gnpa_pct", "GNPA %")):
        xs = [(hist[i] if k == "gnpa_pct" else r[i]).get(k) for i in range(len(hist))]
        prev = [x for x in xs[-4:-1] if x is not None]
        out.append({"key": k, "label": label, "kind": "pct", "year": y, "last": xs[-1],
                    "typical": statistics.median(prev) if prev else None})
    return out


def cagrs(hist, labels_, years):
    """Every rolling `years`-year advances CAGR (model.cagrs' shape)."""
    return model.cagrs([{"revenue": h["advances"]} for h in hist], labels_, years)


# ---------- peers ----------

PEERS = (  # key, label, kind, direction (model.PEERS)
    ("mcap", "Market cap, Rs crore", "num", 1), ("pb", "P/B", "x", -1), ("pe", "P/E", "x", -1),
    ("advances", "Advances, Rs crore", "num", 1), ("g3", "Loan growth, 3 years", "pct", 1),
    ("nim", "NIM", "pct", 1), ("cost_income", "Cost-to-income", "pct", -1), ("credit_cost", "Credit cost", "pct", -1),
    ("roa", "RoA", "pct", 1), ("roe", "RoE", "pct", 1), ("gnpa_pct", "GNPA %", "pct", -1), ("cet1", "CET1", "pct", 1),
)


def peer_metrics(hist, price=None, shares=None):
    """A lender's comparison measures from its last reported year, at a price and share count."""
    h, n, r = hist[-1], len(hist), ratios(hist)[-1]
    mcap = price * shares if price and shares and price > 0 and shares > 0 else None
    return {"mcap": mcap, "pb": mcap / h["equity"] if mcap and h["equity"] > 0 else None,
            "pe": mcap / h["pat"] if mcap and h["pat"] > 0 else None, "advances": h["advances"],
            "g3": model._cagr(hist[-4]["advances"], h["advances"], 3) if n > 3 else None,
            "nim": r["nim"], "cost_income": r["cost_income"], "credit_cost": r["credit_cost"], "roa": r["roa"],
            "roe": r["roe"], "gnpa_pct": h.get("gnpa_pct"), "cet1": h.get("cet1")}


# ---------- Excel ----------

HIST_FORMULA = {  # reported columns: lines derived from the reported ones, as derive() and ratios() do
    "nii": "={int_earned}-{int_expended}", "total_income": "={nii}+{other_income}",
    "equity": "={share_cap}+{reserves}", "total_liab": "={equity}+N({deposits})+{borrowings}+{other_liab}",
    "chk_bs": "=ROUND({total_assets}-{total_liab},2)",
}


def _proj(k, x, base):
    """Projected column formulas; x(key, prev) -> cell ref, base(key) -> the last reported year's cell (absolute)."""
    fund = f"({base('total_assets')}-{base('equity')})"
    return {
        "advances": f"={x('advances', 1)}*(1+{x('adv_growth')})",
        "total_assets": f"={x('advances')}/{x('adv_assets')}",
        "nii": f"={x('nim')}*({x('total_assets', 1)}+{x('total_assets')})/2",
        "other_income": f"={x('oi_pct')}*({x('total_assets', 1)}+{x('total_assets')})/2",
        "total_income": f"={x('nii')}+{x('other_income')}",
        "opex": f"={x('cost_income')}*{x('total_income')}",
        "ppop": f"={x('total_income')}-{x('opex')}",
        "provisions": f"={x('credit_cost')}*({x('advances', 1)}+{x('advances')})/2",
        "pbt": f"={x('ppop')}-{x('provisions')}",
        "tax": f"=MAX(0,{x('pbt')}*{x('tax_rate')})",
        "pat": f"={x('pbt')}-{x('tax')}",
        "dividend": f"=MAX(0,{x('pat')}*{x('payout')})",
        "rwa": f"={x('rwa_density')}*{x('total_assets')}",
        "raised": f"=MAX(0,{x('cet1_min')}*{x('rwa')}-({x('equity', 1)}+{x('pat')}-{x('dividend')}))",
        "share_cap": f"={x('share_cap', 1)}",
        "reserves": f"={x('reserves', 1)}+{x('pat')}-{x('dividend')}+{x('raised')}",
        "equity": f"={x('share_cap')}+{x('reserves')}",
        "cet1": f"=IFERROR(IF({x('rwa')}>0,{x('equity')}/{x('rwa')},\"\"),\"\")",
        "deposits": f"=({x('total_assets')}-{x('equity')})*IFERROR(N({base('deposits')})/{fund},0)",
        "borrowings": f"=({x('total_assets')}-{x('equity')})*IFERROR({base('borrowings')}/{fund},0)",
        "other_liab": f"={x('total_assets')}-{x('equity')}-{x('deposits')}-{x('borrowings')}",
        "total_liab": f"={x('equity')}+{x('deposits')}+{x('borrowings')}+{x('other_liab')}",
        "investments": f"={x('total_assets')}*{base('investments')}/{base('total_assets')}",
        "cash": f"={x('total_assets')}*{base('cash')}/{base('total_assets')}",
        "fixed_assets": f"={x('fixed_assets', 1)}",
        "other_assets": f"={x('total_assets')}-{x('advances')}-{x('investments')}-{x('cash')}-{x('fixed_assets')}",
        "roa": f"=IFERROR({x('pat')}/(({x('total_assets', 1)}+{x('total_assets')})/2),\"\")",
        "roe": f"=IFERROR(IF(({x('equity', 1)}+{x('equity')})/2>1E-6,{x('pat')}/(({x('equity', 1)}+{x('equity')})/2),\"\"),\"\")",
        "leverage": f"=IFERROR(IF(({x('equity', 1)}+{x('equity')})/2>1E-6,({x('total_assets', 1)}+{x('total_assets')})/({x('equity', 1)}+{x('equity')}),\"\"),\"\")",
        "cd_ratio": f"=IFERROR(IF({x('deposits')}>1E-6,{x('advances')}/{x('deposits')},\"\"),\"\")",
        "chk_bs": f"=ROUND({x('total_assets')}-{x('total_liab')},2)",
    }.get(k)


def to_excel(co, scenarios, n=5, active="base", val=None, peers=None, at=None):
    """A lender's workbook: reported years as values, projections as live formulas (inputs blue, picked from the
    Scenarios sheet by the selector in B3), and a Valuation sheet: residual income, P/B, justified P/B, implied
    RoE and the Ke x RoE grid, all formulas. Mirrors project() and valuation()."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = co["name"]
    ws["A1"].data_type = "s"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Rs crore. Lender template: NIM and fee income on average total assets, credit cost on average advances; "
                "equity is raised at year end when CET1 (net worth / RWA) would fall under the floor. Blue = input, "
                "green = picked from the Scenarios sheet, black = formula.")
    ws["A2"].font = Font(italic=True, color="808080")
    ws["A3"] = "Scenario (1 = base, 2 = bull, 3 = bear)"
    ws["B3"] = SCENARIOS.index(active) + 1
    ws["B3"].font = Font(bold=True, color="0000FF")
    ws["C3"] = '=CHOOSE(B3,"Base","Bull","Bear")'
    pick = DataValidation(type="whole", operator="between", formula1="1", formula2="3", showErrorMessage=True,
                          error="Enter 1 (base), 2 (bull) or 3 (bear).")
    ws.add_data_validation(pick)
    pick.add("B3")
    top, nh = 4, len(co["hist"])
    rows = {key: top + 1 + i for i, (_, key, _) in enumerate(LAYOUT) if key}
    headers = labels(co["dates"]) + [fy(d, True) for d in model.projected_dates(co, n)]
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="DDE4EE")
    formats = {"pct": "0.00%", "days": "0", "num": "#,##0.0;(#,##0.0);-", "chk": "#,##0.00;(#,##0.00);-",
               "x": '0.0"x";-0.0"x"'}
    src = model._scenario_sheet(wb, scenarios, n, headers[nh:], formats, head_fill, DRIVERS)
    hist_r = ratios(co["hist"])
    ws.cell(top, 1, "Rs crore").font = bold
    for j, h in enumerate(headers):
        c = ws.cell(top, 2 + j, h)
        c.font = bold
        if j >= nh:
            c.fill = PatternFill("solid", fgColor="FFF2CC")
    last_col = get_column_letter(1 + nh)
    base = lambda k: f"${last_col}${rows[k]}"  # noqa: E731
    totals = {"pat", "total_liab", "total_assets", "equity"}
    for i, (label, key, kind) in enumerate(LAYOUT):
        r = top + 1 + i
        ws.cell(r, 1, label)
        if kind == "head":
            for j in range(len(headers) + 1):
                ws.cell(r, 1 + j).fill = head_fill
            ws.cell(r, 1).font = bold
        if not key:
            continue
        for j in range(len(headers)):
            col, prev = get_column_letter(2 + j), get_column_letter(1 + j)
            x = lambda k, p=0, col=col, prev=prev: f"{prev if p else col}{rows[k]}"  # noqa: E731
            cell = ws.cell(r, 2 + j)
            if j < nh:
                h = co["hist"][j]
                f = HIST_FORMULA.get(key)
                if f and not (key == "total_liab" and h["deposits"] is None):
                    cell.value = f.format(**{k: x(k) for k in ("int_earned", "int_expended", "nii", "other_income",
                                                                "share_cap", "reserves", "equity", "deposits",
                                                                "borrowings", "other_liab", "total_assets", "total_liab")})
                elif key in h:
                    cell.value = h[key]
                elif key in hist_r[j]:
                    cell.value = hist_r[j][key]
                if key == "total_liab" and h["deposits"] is None:
                    cell.value = f"={x('equity')}+{x('borrowings')}+{x('other_liab')}"
            elif key in KEYS:
                cell.value = f"=CHOOSE($B$3,{','.join(src(s, key, j - nh) for s in SCENARIOS)})"
            else:
                cell.value = _proj(key, x, base)
            cell.number_format = formats[kind]
            cell.font = Font(bold=key in totals, color="008000" if j >= nh and key in KEYS else None)
        ws.cell(r, 1).font = Font(bold=key in totals)
    ws.column_dimensions["A"].width = 40
    for j in range(len(headers)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 12
    ws.freeze_panes = ws.cell(top + 1, 2)
    ex = {s: expand(scenarios[s], n) for s in SCENARIOS}
    v = valuation({s: project(co["hist"], ex[s], n) for s in SCENARIOS}, co["hist"], val, co.get("market"))
    _valuation_sheet(wb, co, n, v, lambda k, j: f"Model!{get_column_letter(2 + j)}{rows[k]}", headers[nh:], formats,
                     head_fill)
    if peers:
        model._peers_sheet(wb, peers, at or model.datetime.now(), formats, PEERS)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


VAL_ROWS = ("rf", "beta", "erp", "ke_build", "ke_override", "ke", "g", "price", "shares", "book", "bvps", "nnpa",
            "abvps", "pb", "pabv", "t", "pat", "open", "ri", "df", "pv", "pv_sum", "roe_t", "tv", "pv_tv", "value",
            "per_share", "upside", "justified_pb", "fv_jpb", "implied_roe", "grid")


def _valuation_sheet(wb, co, n, v, at, years, formats, head_fill):
    """'Valuation' sheet: residual income of the scenario Model!B3 picks, as formulas. at(key, j) -> Model ref."""
    inp, nh = v["inputs"], len(co["hist"])
    vs = wb.create_sheet("Valuation")
    vs["A1"] = "Residual income valuation, for the scenario picked in Model!B3"
    vs["A1"].font = Font(bold=True, size=14)
    vs["A2"] = (f"From the {labels(co['dates'])[-1]} book. Value = book + PV of (PAT - Ke x opening book) + PV of "
                "terminal RI = (steady-state RoE - Ke) x final book / (Ke - g). No FCFF DCF: a lender's debt is its "
                "raw material. Blue = input. Rs crore; per share in Rs.")
    vs["A2"].font = Font(italic=True, color="808080")
    blue, bold = Font(color="0000FF"), Font(bold=True)
    ref, r = {}, 4

    def put(key, label, value, fmt="num", font=None):
        nonlocal r
        vs.cell(r, 1, label)
        c = vs.cell(r, 2, value)
        c.number_format = formats.get(fmt, fmt)
        if font:
            c.font = font
        ref[key] = f"$B${r}"
        r += 1

    put("rf", "Risk-free rate", inp["rf"], "pct", blue)
    put("beta", "Beta", inp["beta"], "0.00", blue)
    put("erp", "Equity risk premium", inp["erp"], "pct", blue)
    put("ke_build", "Ke = rf + beta x ERP", f"={ref['rf']}+{ref['beta']}*{ref['erp']}", "pct")
    put("ke_override", "Ke typed (blank = the build-up)", inp["wacc_override"], "pct", blue)
    put("ke", "Ke used", f"=IF(ISNUMBER({ref['ke_override']}),{ref['ke_override']},{ref['ke_build']})", "pct", bold)
    put("g", "Terminal growth", inp["g"], "pct", blue)
    ok = f"AND({ref['ke']}>0,ROUND({ref['ke']}-{ref['g']},10)>={model.MIN_SPREAD})"  # model._clears
    vs.cell(r, 1, f'=IF({ok},"","Ke must be at least 0.5 pp above terminal growth")').font = Font(color="C00000")
    r += 1
    put("price", "CMP (Rs)", v["price"], "#,##0.00", blue)
    put("shares", "Shares (crore)", v["shares"], "#,##0.00", blue)
    put("book", f"Net worth, {labels(co['dates'])[-1]}", f"={at('equity', nh - 1)}")
    has = lambda k: f"ISNUMBER({ref[k]})"  # noqa: E731
    put("bvps", "Book value per share (BVPS)", f"=IF({has('shares')},{ref['book']}/{ref['shares']},\"\")", "#,##0.00")
    put("nnpa", "Net NPA (blank = not filed)", co["hist"][-1].get("nnpa"))
    put("abvps", "Adjusted BVPS (book less net NPA)",
        f"=IF(AND({has('shares')},{has('nnpa')}),({ref['book']}-{ref['nnpa']})/{ref['shares']},\"\")", "#,##0.00")
    put("pb", "P/B at CMP", f"=IFERROR(IF(AND({has('price')},{ref['bvps']}>0),{ref['price']}/{ref['bvps']},\"\"),\"\")", "x")
    put("pabv", "P/ABV at CMP", f"=IFERROR(IF(AND({has('price')},{ref['abvps']}>0),{ref['price']}/{ref['abvps']},\"\"),\"\")", "x")
    r += 1
    vs.cell(r, 1, "Rs crore").font = bold
    for j, y in enumerate(years):
        vs.cell(r, 2 + j, y).font = bold
        vs.cell(r, 2 + j).fill = head_fill
    r += 1
    rows = {}
    for key, label in (("t", "Year (t)"), ("pat", "Net profit"), ("open", "Opening book"), ("ri", "Residual income"),
                       ("df", "Discount factor"), ("pv", "PV of residual income")):
        vs.cell(r, 1, label)
        rows[key] = r
        for j in range(n):
            c, cl = vs.cell(r, 2 + j), get_column_letter(2 + j)
            c.value = {"t": j + 1, "pat": f"={at('pat', nh + j)}",
                       "open": f"={at('equity', nh + j - 1)}",
                       "ri": f"={cl}{rows.get('pat')}-{ref['ke']}*{cl}{rows.get('open')}",
                       "df": f"=1/(1+{ref['ke']})^{cl}{rows['t']}",
                       "pv": f"={cl}{rows.get('ri')}*{cl}{rows.get('df')}"}[key]
            c.number_format = "0" if key == "t" else "0.0000" if key == "df" else formats["num"]
        r += 1
    rng = lambda k: f"$B${rows[k]}:${get_column_letter(1 + n)}${rows[k]}"  # noqa: E731
    lastc = get_column_letter(1 + n)
    put("pv_sum", "Sum of PV of residual income", f"=IF({ok},SUM({rng('pv')}),\"\")")
    put("roe_t", "Steady-state RoE (last year's PAT / opening book)",
        f"=IFERROR({lastc}{rows['pat']}/{lastc}{rows['open']},\"\")", "pct")
    put("final", "Final-year book", f"={at('equity', nh + n - 1)}")
    put("tv", "Terminal value of residual income",
        f"=IF({ok},({ref['roe_t']}-{ref['ke']})*{ref['final']}/({ref['ke']}-{ref['g']}),\"\")")
    put("pv_tv", "PV of terminal value", f"=IF({ok},{ref['tv']}*{lastc}{rows['df']},\"\")")
    put("value", "Equity value = book + PV of RI + PV of TV", f"=IF({ok},{ref['book']}+{ref['pv_sum']}+{ref['pv_tv']},\"\")",
        "num", bold)
    put("per_share", "Fair value per share (Rs)", f"=IFERROR(IF(AND({ok},{has('shares')}),{ref['value']}/{ref['shares']},\"\"),\"\")",
        "#,##0.00", bold)
    put("upside", "Upside / (downside) to CMP",
        f"=IFERROR(IF(AND({has('price')},ISNUMBER({ref['per_share']})),{ref['per_share']}/{ref['price']}-1,\"\"),\"\")", "pct")
    put("justified_pb", "Justified P/B = (RoE - g) / (Ke - g)",
        f"=IFERROR(IF({ok},({ref['roe_t']}-{ref['g']})/({ref['ke']}-{ref['g']}),\"\"),\"\")", "x")
    put("fv_jpb", "Fair value at justified P/B (Rs)",
        f"=IFERROR({ref['justified_pb']}*{ref['bvps']},\"\")", "#,##0.00")
    put("implied_roe", "Market-implied steady-state RoE = g + P/B x (Ke - g)",
        f"=IFERROR(IF(AND({ok},ISNUMBER({ref['pb']})),{ref['g']}+{ref['pb']}*({ref['ke']}-{ref['g']}),\"\"),\"\")", "pct")
    r += 1
    vs.cell(r, 1, "Fair value per share (Rs): Ke (rows) x steady-state RoE (columns)").font = bold
    r += 1
    for j, dr in enumerate(GRID_ROE):
        c = vs.cell(r, 2 + j, f"={ref['roe_t']}+{dr}")
        c.number_format, c.font = formats["pct"], bold
    grid_top = r
    for i, dk in enumerate(GRID_KE):
        rr = grid_top + 1 + i
        k = f"$A{rr}"
        a = vs.cell(rr, 1, f"={ref['ke']}+{dk}")
        a.number_format, a.font = formats["pct"], bold
        for j in range(len(GRID_ROE)):
            roe = f"{get_column_letter(2 + j)}${grid_top}"
            body = (f"({ref['book']}+SUMPRODUCT(({rng('pat')}-{k}*{rng('open')})/(1+{k})^{rng('t')})"
                    f"+({roe}-{k})*{ref['final']}/({k}-{ref['g']})/(1+{k})^{n})/{ref['shares']}")
            c = vs.cell(rr, 2 + j, f"=IFERROR(IF(AND({k}>0,ROUND({k}-{ref['g']},10)>={model.MIN_SPREAD},{has('shares')},"
                                   f"ISNUMBER({roe})),{body},\"\"),\"\")")
            c.number_format = "#,##0.00"
    ref["grid"] = (grid_top + 1, 2)
    vs.column_dimensions["A"].width = 52
    for j in range(max(n, len(GRID_ROE))):
        vs.column_dimensions[get_column_letter(2 + j)].width = 13
    vs.sheet_properties.tabColor = "DDE4EE"
    return ref
